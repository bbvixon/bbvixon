from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = r"D:\agent\bbvixon\may-be-design-system.docx"

BLACK = "000000"
INK = "332119"
BROWN = "3A241B"
MUTED = "6F625A"
CREAM = "F5F0E7"
SURFACE = "FFFAF2"
LINE = "D8CEC5"
PALE = "F8F5F0"
GREEN = "3F6B4F"
AMBER = "8A5A18"
RED = "B3261E"


def set_font(run, size=None, bold=None, color=INK, name="Segoe UI", italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def margins(cell, top=150, start=150, bottom=150, end=150):
    tcpr = cell._tc.get_or_add_tcPr()
    tcmar = tcpr.first_child_found_in("w:tcMar")
    if tcmar is None:
        tcmar = OxmlElement("w:tcMar")
        tcpr.append(tcmar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcmar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tcmar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def borders(cell, color=LINE, size="5"):
    tcpr = cell._tc.get_or_add_tcPr()
    box = tcpr.first_child_found_in("w:tcBorders")
    if box is None:
        box = OxmlElement("w:tcBorders")
        tcpr.append(box)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = box.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            box.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def remove_paragraph_borders(paragraph_or_style):
    ppr = paragraph_or_style._element.get_or_add_pPr()
    box = ppr.find(qn("w:pBdr"))
    if box is not None:
        ppr.remove(box)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    trpr.append(node)


def eyebrow(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text.upper())
    set_font(r, size=9, bold=True, color=MUTED)
    return p


def section_title(doc, number, title, intro=None):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = True
    r = p.add_run(title)
    set_font(r, size=28, bold=True, color=BLACK)
    if intro:
        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(22)
        r2 = p2.add_run(intro)
        set_font(r2, size=12.5, color=MUTED)


def bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("-  " + text)
    set_font(r, size=11.5)
    return p


def label(doc, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(title + "  ")
    set_font(r, size=11.5, bold=True, color=BLACK)
    r = p.add_run(text)
    set_font(r, size=11.5, color=INK)
    return p


def table(doc, headers, rows, widths, header_fill=BROWN, swatches=None, bold_col=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    hdr = t.rows[0]
    repeat_header(hdr)
    for i, text in enumerate(headers):
        c = hdr.cells[i]
        c.width = Inches(widths[i])
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        margins(c, 140, 150, 140, 150)
        borders(c)
        shade(c, header_fill)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_font(r, size=10, bold=True, color="FFFFFF")
    for row_i, row in enumerate(rows):
        cells = t.add_row().cells
        for col_i, text in enumerate(row):
            c = cells[col_i]
            c.width = Inches(widths[col_i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            margins(c, 135, 150, 135, 150)
            borders(c)
            if row_i % 2:
                shade(c, PALE)
            if swatches and col_i == 0:
                shade(c, swatches[row_i])
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(text))
            text_color = "FFFFFF" if swatches and col_i == 0 and swatches[row_i] in (INK, BROWN, GREEN, AMBER, RED) else INK
            set_font(r, size=10.5, bold=(col_i == bold_col), color=text_color)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def add_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("may.be   Design System")
    set_font(r, size=8.5, color=MUTED)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.82)
sec.bottom_margin = Inches(0.75)
sec.left_margin = Inches(0.9)
sec.right_margin = Inches(0.9)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Segoe UI"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Segoe UI")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Segoe UI")
normal.font.size = Pt(11.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.16

for name, size, before, after in [
    ("Title", 34, 0, 10),
    ("Subtitle", 15, 0, 20),
    ("Heading 1", 28, 0, 16),
    ("Heading 2", 17, 17, 7),
]:
    s = styles[name]
    s.font.name = "Segoe UI"
    s._element.rPr.rFonts.set(qn("w:ascii"), "Segoe UI")
    s._element.rPr.rFonts.set(qn("w:hAnsi"), "Segoe UI")
    s.font.size = Pt(size)
    s.font.color.rgb = RGBColor.from_string(BLACK)
    if name != "Subtitle":
        s.font.bold = True
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
    remove_paragraph_borders(s)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(100)
p.paragraph_format.space_after = Pt(28)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("may.be")
set_font(r, size=44, bold=True, color=BROWN)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
remove_paragraph_borders(p)
r = p.add_run("Design System")
set_font(r, size=34, bold=True, color=BLACK)

p = doc.add_paragraph(style="Subtitle")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("A quiet mobile interface for exploring decisions")
set_font(r, size=15, color=MUTED)

doc.add_paragraph().paragraph_format.space_after = Pt(28)

table(doc, ["Clear", "Calm", "Continuous"], [
    ("One purpose per screen", "Space is part of the experience", "Every choice grows from the same tree"),
], [2.15, 2.15, 2.15])

doc.add_paragraph().paragraph_format.space_after = Pt(18)
label(doc, "Platform", "iPhone first. Responsive web and Android use the same semantic system.")
label(doc, "Standard", "Apple Human Interface Guidelines for hierarchy, motion, accessibility, safe areas, and touch behavior.")
label(doc, "Brand", "Warm cream, deep brown, restrained language, and one cinematic focal point.")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(46)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Version 1.0   October 2026")
set_font(r, size=9.5, color=MUTED)

# Page 2
section_title(doc, "01", "Principles", "Apple familiar. may.be distinct.")

for title, text in [
    ("One focus", "Each scene has one primary action and one visual centre."),
    ("Direct manipulation", "The object a person chooses is the object that moves and grows."),
    ("Agency", "Undo is immediate. Exploration never feels like commitment."),
    ("Restraint", "No dashboard cards, decorative gradients, or unnecessary copy."),
    ("Adaptation", "Safe areas, Dynamic Type, appearance modes, and input methods shape the layout."),
    ("Trust", "Plausible paths are never presented as predictions."),
]:
    label(doc, title, text)

doc.add_heading("Screen rhythm", level=2)
table(doc, ["Tell", "Understand", "Transform", "Explore"], [
    ("Speak or type", "Keywords emerge", "Dot becomes decision", "Branches grow"),
], [1.6, 1.6, 1.6, 1.6])

doc.add_heading("Keep", level=2)
bullet(doc, "Native behavior for keyboard, microphone, focus, permissions, and safe areas.")
bullet(doc, "may.be styling in the content layer: color, graph, pacing, and language.")

# Page 3
section_title(doc, "02", "Color", "Warm neutrals carry the story. Brown marks the active decision.")

color_rows = [
    ("", "Canvas", "#F5F0E7", "Atmosphere"),
    ("", "Surface", "#FFFAF2", "Nodes and sheets"),
    ("", "Primary text", "#332119", "Essential content"),
    ("", "Secondary text", "#6F625A", "Supporting content"),
    ("", "Primary action", "#3A241B", "One active control"),
    ("", "Separator", "#D8CEC5", "Borders and paths"),
    ("", "Success", "#3F6B4F", "Confirmed state"),
    ("", "Warning", "#8A5A18", "Recoverable issue"),
    ("", "Destructive", "#B3261E", "Destructive action"),
]
table(doc, ["Swatch", "Token", "Light", "Use"], color_rows, [0.75, 1.65, 1.2, 2.8], swatches=[CREAM, SURFACE, INK, MUTED, BROWN, LINE, GREEN, AMBER, RED], bold_col=1)

doc.add_heading("Rules", level=2)
bullet(doc, "Use semantic color names in Figma and code. Do not name tokens by hex value.")
bullet(doc, "Reserve deep brown for the decision seed, selected path, or one primary action.")
bullet(doc, "Provide dark and Increased Contrast variants. Prefer system semantic colors for standard controls.")
bullet(doc, "Never use color alone. Pair state with text, icon, line weight, or motion.")

label(doc, "Contrast", "4.5 to 1 for normal text. 3 to 1 for large text and essential geometry.")

# Page 4
section_title(doc, "03", "Type and layout", "Use the system font. Let content breathe.")

table(doc, ["Style", "Size", "Weight", "Use"], [
    ("Large title", "34 pt", "Semibold", "Rare cinematic statement"),
    ("Title", "28 pt", "Semibold", "Decision framing"),
    ("Title 3", "20 pt", "Semibold", "Decision seed"),
    ("Body", "17 pt", "Regular", "Input and explanations"),
    ("Callout", "16 pt", "Medium", "Branch labels"),
    ("Subheadline", "15 pt", "Regular", "Secondary instruction"),
    ("Caption", "12 pt", "Regular", "Noncritical hint"),
], [1.35, 1.0, 1.25, 3.05], bold_col=1)

label(doc, "Font", "SF Pro through the Apple system font. Web stack: -apple-system, BlinkMacSystemFont, Segoe UI, sans-serif.")
label(doc, "Case", "Sentence case. Short labels. No all caps in the product interface.")
label(doc, "Scaling", "Support Dynamic Type. Wrap before shrinking. At accessibility sizes, switch the graph to a list presentation.")

doc.add_heading("Spacing", level=2)
table(doc, ["4", "8", "12", "16", "20", "24", "32", "48"], [
    ("Micro", "Inside", "Compact", "Default", "Margin", "Group", "Section", "Pause"),
], [0.8, 0.8, 0.8, 0.8, 0.8, 0.8, 0.8, 0.8])

label(doc, "Frame", "390 by 844 points. 20 point side margin. 44 by 44 point minimum targets.")
label(doc, "Insets", "Respect the system safe area and keyboard inset.")

# Page 5
section_title(doc, "04", "Components", "Few components. Complete states.")

table(doc, ["Component", "Default", "States"], [
    ("Microphone orb", "104 pt circle", "Idle, pressed, listening, unavailable"),
    ("Text composer", "50 pt minimum height", "Empty, focused, submitting, error"),
    ("Keyword bubble", "Content-sized capsule", "Enter, settle, dissolve"),
    ("Decision seed", "120 pt circle", "Growing, current, loading, error"),
    ("Branch node", "96 pt circle", "Resting, focused, selected, rejected"),
    ("Branch path", "1.5 pt line", "Drawing, active, historical"),
    ("Detail caption", "Three lines maximum", "Hint, summary, loading, error"),
    ("Undo", "Symbol and label", "Hidden, available, pressed"),
], [1.7, 1.8, 2.95])

doc.add_heading("Shape", level=2)
label(doc, "Corners", "12 pt controls. 18 pt fields. 24 pt sheets. Full circle for bubbles and nodes.")
label(doc, "Depth", "Use borders and tonal contrast first. Use one restrained shadow only on the active seed.")
label(doc, "Material", "Use platform material for floating navigation only. Keep the graph surface warm and opaque.")

doc.add_heading("Logo rule", level=2)
bullet(doc, "Keep the cinematic dot mathematically centred.")
bullet(doc, "Position may and be around it; word width must never move the dot.")
bullet(doc, "Keep the header logo conventional while the cinematic logo transforms.")

# Page 6
section_title(doc, "05", "Motion", "Motion shows cause and continuity.")

table(doc, ["Moment", "Timing", "Behavior"], [
    ("Press", "100 to 160 ms", "Scale or tone responds immediately"),
    ("Caption", "180 to 260 ms", "Short crossfade"),
    ("Keyword", "420 to 700 ms", "Rise, settle, pause, dissolve"),
    ("Branch", "500 to 800 ms", "Path draws before node arrives"),
    ("Camera", "600 to 900 ms", "Pan and scale one world"),
    ("Logo to seed", "900 to 1400 ms", "Words fade; pinned dot grows"),
], [1.55, 1.45, 3.45])

label(doc, "Easing", "Authored motion: cubic-bezier(.22, 1, .36, 1). Direct touch: a tight spring with minimal overshoot.")

doc.add_heading("Branch sequence", level=2)
table(doc, ["1", "2", "3", "4"], [
    ("Draw path", "Reveal node", "Enable touch", "Prefetch nearby paths"),
], [1.6, 1.6, 1.6, 1.6])

doc.add_heading("Reduced Motion", level=2)
bullet(doc, "Replace zoom, large travel, parallax, and blur with short fades.")
bullet(doc, "Keep the same order, labels, focus, and final state.")
bullet(doc, "Do not remove feedback or meaning.")

# Page 7
section_title(doc, "06", "Accessibility and handoff", "The system is complete only when every path remains usable.")

table(doc, ["Setting", "Required response"], [
    ("VoiceOver", "Decision is a heading. Branches are buttons with label and summary."),
    ("Dynamic Type", "Containers grow. Graph can become a structured list."),
    ("Increase Contrast", "Strengthen text, paths, borders, and focus rings."),
    ("Reduce Transparency", "Use opaque warm surfaces instead of blur."),
    ("Switch or keyboard", "Logical focus order. Enter and Space activate."),
    ("Touch", "44 point minimum targets with clear pressed feedback."),
], [1.75, 4.65])

doc.add_heading("Shared naming", level=2)
label(doc, "Figma", "Color, Type, Spacing, Radius, Motion, and Components collections.")
label(doc, "States", "idle, focused, pressed, listening, processing, current, selected, rejected, loading, error, disabled.")
label(doc, "Code", "Use the same semantic names in SwiftUI and CSS. Each motion token includes a reduced-motion alternative.")

doc.add_heading("Release check", level=2)
bullet(doc, "Test light, dark, Increased Contrast, Reduce Transparency, and Reduce Motion.")
bullet(doc, "Test smallest iPhone width and one accessibility text size.")
bullet(doc, "Confirm every branch stays inside the viewport and Undo restores the exact previous fork.")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(18)
r = p.add_run("Apple references  ")
set_font(r, size=9.5, bold=True, color=BLACK)
r = p.add_run("Design Principles  /  Color  /  Layout  /  Motion  /  Accessibility  /  Typography")
set_font(r, size=9.5, color=MUTED)

for section in doc.sections:
    add_footer(section)

doc.core_properties.title = "may.be Design System"
doc.core_properties.subject = "Apple aligned mobile visual and interaction system"
doc.core_properties.author = "may.be product team"
doc.core_properties.keywords = "may.be, design system, Apple HIG, mobile, color, typography, motion, accessibility"

doc.save(OUT)
print(OUT)
