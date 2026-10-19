from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK


OUT = r"D:\agent\bbvixon\may-be-product-and-design-specification.docx"

INK = "332119"
MUTED = "6F625A"
CREAM = "F5F0E7"
IVORY = "FFFAF2"
ACCENT = "B47B54"
BLACK = "000000"
GRID = "D9D9D9"
PALE = "F7F3ED"
CTA = "3A241B"
CTA_HOVER = "4B3025"
CTA_PRESSED = "281710"
SEPARATOR = "D8CEC5"
SUCCESS = "3F6B4F"
WARNING = "8A5A18"
DESTRUCTIVE = "B3261E"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=GRID, size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Arial", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def remove_paragraph_borders(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Inches(0.25 + 0.2 * level)
    p.paragraph_format.first_line_indent = Inches(-0.16)
    r = p.add_run("- ")
    set_font(r, size=9.5, color=INK)
    r = p.add_run(text)
    set_font(r, size=9.5, color=INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, size=9.5, color=INK)
    return p


def add_label_paragraph(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(label + " ")
    set_font(r, size=9.5, bold=True, color=BLACK)
    r2 = p.add_run(text)
    set_font(r2, size=9.5, color=INK)
    return p


def add_story(doc, story_id, title, statement, criteria, notes=None):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    set_keep_with_next(p)
    r = p.add_run(f"{story_id}  {title}")
    set_font(r, size=12, bold=True, color=BLACK)
    r.add_break()
    label_run = p.add_run("User story ")
    set_font(label_run, size=9.5, bold=True, color=BLACK)
    statement_run = p.add_run(statement)
    set_font(statement_run, size=9.5, color=INK)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(3)
    r2 = p2.add_run("Acceptance criteria")
    set_font(r2, size=9.5, bold=True, color=BLACK)
    set_keep_with_next(p2)
    for item in criteria:
        add_bullet(doc, item)
    if notes:
        add_label_paragraph(doc, "Implementation note", notes)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, INK)
        set_cell_margins(cell, 115, 120, 115, 120)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_font(r, size=8.5, bold=True, color="FFFFFF")
        if widths:
            cell.width = Inches(widths[i])
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            if idx % 2 == 1:
                set_cell_shading(cell, PALE)
            set_cell_margins(cell, 105, 120, 105, 120)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(text))
            set_font(r, size=8.5, color=INK)
            if widths:
                cell.width = Inches(widths[i])
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

# Base typography and styles
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(9.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

for style_name, size, before, after in [
    ("Title", 28, 0, 12),
    ("Subtitle", 12, 0, 18),
    ("Heading 1", 19, 18, 8),
    ("Heading 2", 14, 13, 6),
    ("Heading 3", 12, 10, 4),
]:
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(BLACK)
    if style_name != "Subtitle":
        style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
    remove_paragraph_borders(style)

if "Story Statement" not in styles:
    story_style = styles.add_style("Story Statement", WD_STYLE_TYPE.PARAGRAPH)
    story_style.font.name = "Arial"
    story_style.font.size = Pt(11)
    story_style.font.color.rgb = RGBColor.from_string(INK)
    story_style.paragraph_format.left_indent = Inches(0.18)
    story_style.paragraph_format.right_indent = Inches(0.18)
    story_style.paragraph_format.space_after = Pt(7)

# Cover
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
remove_paragraph_borders(p)
r = p.add_run("may.be Product and Design Specification")
set_font(r, size=28, bold=True, color=BLACK)

p = doc.add_paragraph(style="Subtitle")
r = p.add_run("User journey, Apple aligned design system, interaction model, and engineering contract")
set_font(r, size=12, color=MUTED)

add_label_paragraph(doc, "Purpose", "Define the intended user experience, visual system, interaction behavior, motion language, accessibility requirements, engineering contract, and product boundaries for the first complete may.be journey.")
add_label_paragraph(doc, "Audience", "Product design, frontend engineering, backend engineering, AI integration, quality assurance, and accessibility review.")
add_label_paragraph(doc, "Product promise", "Help a person see plausible paths through a difficult decision without pretending to predict the future or compressing the experience into a dashboard.")

doc.add_paragraph()
flow = add_table(
    doc,
    ["Input", "Interpretation", "Transformation", "Exploration", "Recovery"],
    [["Speak or type", "Keywords and factors", "Story to decision seed", "Living branch graph", "Undo or restart"]],
    [1.18, 1.34, 1.42, 1.30, 1.16],
)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
r = p.add_run("Document status  Product and engineering handoff")
set_font(r, size=9, color=MUTED)

# Overview
doc.add_heading("1 Product intent", level=1)
doc.add_paragraph(
    "may.be is a mobile first decision exploration experience. A person describes where they are stuck, the system extracts grounded meaning, distils the central decision, and reveals a spatial tree of credible alternatives. The graph remains exploratory: choosing a branch means following it to understand what it could lead to, not committing to it."
)
doc.add_paragraph(
    "The primary product requirement is continuity. The user should feel that one story becomes one decision seed and then grows into one persistent tree. Screens may change, but the experience must not feel like a sequence of unrelated cards, reports, or dashboard panels."
)

doc.add_heading("1.1 Core principles", level=2)
for item in [
    "Voice is the primary input; typing is a quiet and fully functional fallback.",
    "AI interpretation stays visibly connected to the words the person supplied.",
    "The graph represents plausible paths, not predicted or guaranteed outcomes.",
    "Motion explains transformation and continuity, so interface copy remains minimal.",
    "The number of branches is determined by distinct credible options, not a fixed layout quota.",
    "Every exploration action is reversible without refreshing or rebuilding the page.",
    "Human unknowns and assumptions remain distinguishable from user facts and evidence.",
]:
    add_bullet(doc, item)

doc.add_heading("1.2 Primary persona", level=2)
add_label_paragraph(doc, "Person", "An adult facing a meaningful personal, career, financial, relationship, or product choice who feels overwhelmed by competing consequences.")
add_label_paragraph(doc, "Need", "A calm way to externalise the decision, see distinct alternatives, and follow consequences without reading a long AI answer.")
add_label_paragraph(doc, "Success", "The person can articulate the decision more clearly, understand the credible paths, inspect why each path exists, and return to earlier forks safely.")

doc.add_heading("1.3 Scope of the first complete journey", level=2)
add_table(
    doc,
    ["Included", "Deferred"],
    [["Voice and text capture", "Accounts and cross device sync"],
     ["Local AI analysis and branch expansion", "Evidence retrieval with citations"],
     ["Cinematic keyword and logo transition", "Probability or outcome scoring"],
     ["Variable branch graph with details", "Autonomous multi agent planning"],
     ["Selection, expansion, undo, restart", "Saved Your paths constellation"]],
    [3.15, 3.15],
)

# Design system
doc.add_heading("2 Apple aligned design system", level=1)
doc.add_paragraph(
    "may.be follows Apple Human Interface Guidelines at the level of behavior, hierarchy, legibility, safe-area awareness, touch ergonomics, accessibility, and purposeful motion. It keeps its own warm, contemplative identity rather than imitating Apple screens. Native controls and semantic platform colors should be used when they improve familiarity; custom graph and cinematic elements must still behave like trustworthy iOS interactions."
)

doc.add_heading("2.1 Design character", level=2)
add_table(doc, ["Attribute", "Design decision", "Avoid"], [
    ("Calm", "One focal action per scene, generous negative space, slow state changes", "Dense cards, dashboards, competing controls"),
    ("Grounded", "Warm cream surfaces and dark brown interactive anchors", "Cold blue AI styling or neon gradients"),
    ("Cinematic", "Transformation preserves object identity and spatial continuity", "Hard cuts, unrelated screen replacements, decorative motion"),
    ("Human", "Brief direct language and visibly grounded AI interpretation", "Robotic analysis labels or invented certainty"),
    ("Native", "Safe areas, 44 point targets, Dynamic Type, standard gestures and focus", "Web-only patterns that feel foreign on iPhone"),
], [1.05, 3.25, 2.05])

doc.add_heading("2.2 Color theory", level=2)
doc.add_paragraph(
    "The palette uses a warm cream field to reduce the clinical feeling common in AI products. Deep brown carries weight without the aggression of pure black and becomes the visual metaphor for a decision seed: grounded, serious, and capable of growth. Most of the interface remains chromatically quiet so motion and geometry can communicate structure. Color never carries meaning alone, and accent color is reserved for an active choice or primary action."
)
add_table(doc, ["Semantic token", "Light value", "Dark value", "Role"], [
    ("background.canvas", "#F5F0E7", "#171310", "Full-screen atmosphere behind the phone canvas"),
    ("background.surface", "#FFFAF2", "#211A16", "Primary app surface and node interiors"),
    ("text.primary", "#332119", "#F7EFE8", "Headings, labels, and essential content"),
    ("text.secondary", "#6F625A", "#C8BBB1", "Supporting copy; never primary instructions"),
    ("action.primary", "#3A241B", "#D5A27F", "Microphone, decision seed, and selected path"),
    ("action.hover", "#4B3025", "#E0B18F", "Pointer hover or highlighted state"),
    ("action.pressed", "#281710", "#BE8968", "Pressed state with scale or haptic feedback"),
    ("border.subtle", "#D8CEC5", "#4A3D36", "Hairlines, node outlines, and separators"),
    ("status.success", "#3F6B4F", "#78B68C", "Confirmed completion, paired with icon or text"),
    ("status.warning", "#8A5A18", "#E1AD5A", "Recoverable warning, paired with explanation"),
    ("status.destructive", "#B3261E", "#FF8A80", "Destructive confirmation only"),
], [1.35, 0.95, 0.95, 3.1])
add_label_paragraph(doc, "Contrast rule", "Normal text must meet at least 4.5 to 1 against its resting background; large text and essential geometry must meet at least 3 to 1. The current primary brown on cream exceeds 12 to 1, while secondary text uses the darker #6F625A to remain above 4.5 to 1 on both warm surfaces.")
add_label_paragraph(doc, "Usage rule", "Do not fill several neighboring controls with the action color. The strongest brown belongs to the single current decision or primary action. Unselected branches use the surface color with an outline so the hierarchy remains obvious.")

doc.add_heading("2.3 Typography", level=2)
doc.add_paragraph(
    "Use San Francisco through the system font on Apple platforms. On the web, use the system stack so an iPhone resolves naturally to SF Pro while Windows and Android receive an appropriate native substitute. Do not package or redistribute Apple font files for non-Apple platforms. Support Dynamic Type and allow labels to wrap before shrinking them."
)
add_table(doc, ["Text style", "iOS reference", "Weight", "may.be use"], [
    ("Large title", "34 pt", "Semibold", "Rare transition statement; maximum two lines"),
    ("Title 1", "28 pt", "Semibold", "Decision framing and major confirmation"),
    ("Title 3", "20 pt", "Semibold", "Decision seed when space permits"),
    ("Body", "17 pt", "Regular", "Input, path detail, errors, and explanations"),
    ("Callout", "16 pt", "Regular or medium", "Branch-node labels"),
    ("Subheadline", "15 pt", "Regular", "Secondary instructions and status"),
    ("Caption 1", "12 pt", "Regular", "Transient hints only; never critical content"),
], [1.25, 1.15, 1.25, 3.0])
add_label_paragraph(doc, "Font stack", "-apple-system, BlinkMacSystemFont, SF Pro Text, Segoe UI, sans-serif. Use SF Pro Display behavior for titles at 20 points and above and Text behavior below that threshold.")
add_label_paragraph(doc, "Text behavior", "Use sentence case, natural line breaks, and left alignment for paragraphs. Centre only cinematic statements, nodes, and very short labels. Avoid all caps, justified copy, and more than two weights in one scene.")

doc.add_heading("2.4 Layout and spacing", level=2)
add_table(doc, ["Token or rule", "Value", "Application"], [
    ("space.1", "4 pt", "Icon-to-label micro spacing"),
    ("space.2", "8 pt", "Related elements inside a control"),
    ("space.3", "12 pt", "Compact component padding"),
    ("space.4", "16 pt", "Default vertical rhythm"),
    ("space.5", "20 pt", "Minimum phone side margin"),
    ("space.6", "24 pt", "Component-group separation"),
    ("space.8", "32 pt", "Major scene spacing"),
    ("space.12", "48 pt", "Cinematic breathing space"),
    ("Touch target", "44 by 44 pt minimum", "Every tappable node and control"),
    ("Safe area", "System provided", "No control or caption beneath sensor housing or home indicator"),
], [1.45, 1.3, 3.9])
doc.add_paragraph(
    "The primary phone reference frame is 390 by 844 points, but the layout is fluid from compact iPhone widths upward. The microphone and decision seed sit on the horizontal centreline. The bottom input follows the keyboard and safe-area inset. The graph occupies a world coordinate system that can pan and zoom within a clipped viewport; branch labels never determine absolute screen coordinates."
)

doc.add_heading("2.5 Shape, depth, and material", level=2)
for item in [
    "Use continuous rounded corners where the platform supports them. Suggested radii are 12 points for compact controls, 18 points for fields and captions, 24 points for sheets, and a true circle for bubbles and graph nodes.",
    "Prefer borders, tonal separation, and restrained ambient shadow over stacked cards. The decision graph is the content, not a card inside another card.",
    "Use system materials only for transient navigation or controls that float above moving content. Keep the central story and graph surfaces warm and mostly opaque.",
    "If Reduce Transparency is active, replace blur and translucent material with an opaque surface and a clear separator.",
    "Shadow tokens: node resting 0 8 24 at 8 percent brown; selected seed 0 12 36 at 14 percent brown. Shadows must not be the sole boundary cue.",
]:
    add_bullet(doc, item)

doc.add_heading("2.6 Component specifications", level=2)
add_table(doc, ["Component", "Anatomy and size", "States and behavior"], [
    ("Brand mark", "may, centred dot, be; small header variant", "Header remains stable; cinematic mark permits word halves to fade while dot stays pinned"),
    ("Microphone orb", "96 to 120 pt circle; SF Symbol mic; 44 pt hit minimum", "Idle, pressed, listening, processing, unavailable; haptic on start and stop"),
    ("Text composer", "Minimum 50 pt height; 18 pt radius; send control inside trailing edge", "Empty, focused, populated, submitting, error; follows keyboard safe area"),
    ("Keyword bubble", "Content-sized capsule; one or two short words", "Emerges, settles, dissolves; never tappable unless correction is introduced"),
    ("Decision seed", "112 to 136 pt circle with two-line maximum label", "Grows from logo dot, current, selected, loading, error"),
    ("Branch node", "88 to 104 pt circle; short label; 44 pt minimum hit area", "Resting, hover or focus, pressed, selected, rejected, loading"),
    ("Branch path", "1.5 pt default; 2 pt active; rounded caps", "Draws from source, highlights with selection, remains spatially attached"),
    ("Detail caption", "Fixed bottom region above safe area; maximum three short lines", "Neutral hint, branch summary, loading, recoverable error"),
    ("Undo", "System chevron-back or undo symbol plus text", "Appears only with history; reverses one graph step without page refresh"),
], [1.25, 2.45, 2.95])

doc.add_heading("2.7 Motion system", level=2)
doc.add_paragraph(
    "Motion explains causality: words become meaning, the logo dot becomes the decision, and a selected branch becomes the next centre. The same object should persist whenever the concept persists. Use spring behavior for direct manipulation and a smooth cubic easing for authored cinematic transitions. Motion must remain interruptible where the user controls it."
)
add_table(doc, ["Motion token", "Duration", "Use"], [
    ("instant", "100 to 160 ms", "Pressed feedback, focus response, icon state"),
    ("responsive", "180 to 260 ms", "Small control transitions and caption changes"),
    ("spatial", "420 to 700 ms", "Branch growth, sibling rejection, camera pan"),
    ("cinematic", "800 to 1400 ms", "Passage copy, wordmark, seed transformation"),
    ("breathing hold", "600 to 1200 ms", "Intentional pause between emotional beats"),
], [1.4, 1.3, 3.9])
add_label_paragraph(doc, "Primary easing", "cubic-bezier(.22, 1, .36, 1) for authored movement. Direct touch may use a critically damped spring with minimal overshoot.")
add_label_paragraph(doc, "Reduced motion", "Remove scale travel, camera fly-through, parallax, repeated breathing, and blur transitions. Replace them with short opacity crossfades while keeping the same semantic order. Never remove status or information.")

doc.add_heading("2.8 Interaction and feedback", level=2)
for item in [
    "Primary actions respond on touch-down with subtle scale or tonal feedback and complete on touch-up inside the target.",
    "Hover is enhancement only. Touch reveals branch detail on first interaction and follows the branch through an explicit tap or a clearly staged second action where accidental commitment is likely.",
    "Use light haptics when listening starts, a soft success haptic when analysis completes, and selection feedback when a branch becomes current. Do not vibrate for every animation beat.",
    "Disable duplicate submission and branch selection during state transitions, but preserve visible feedback so the interface never appears frozen.",
    "Destructive restart requires confirmation only after meaningful user work exists; Undo remains immediate because it is reversible.",
]:
    add_bullet(doc, item)

doc.add_heading("2.9 Accessibility", level=2)
add_table(doc, ["Need", "Requirement"], [
    ("VoiceOver", "Expose the decision as a heading and every branch as a button whose label combines the short title and concise explanation. Announce new levels politely."),
    ("Dynamic Type", "Support system scaling without clipping. At accessibility sizes, replace circular text layout with an accessible list while preserving the graph relationship."),
    ("Reduce Motion", "Use fades instead of large translations, zooms, depth changes, or blur. Preserve sequence and focus."),
    ("Increase Contrast", "Strengthen node borders, paths, focus rings, and secondary text; never depend on soft shadows."),
    ("Reduce Transparency", "Replace materials and background blur with opaque warm surfaces."),
    ("Differentiate Without Color", "Use labels, line weight, icons, focus rings, and shape or motion states in addition to color."),
    ("Switch and keyboard", "Maintain logical focus order: header, current decision, branches from left to right, caption, Undo. Enter and Space activate."),
], [1.65, 5.0])

doc.add_heading("2.10 Content design", level=2)
for item in [
    "Opening copy: Hi there. Tell me where you are stuck. Keep the screen otherwise quiet.",
    "Use verbs for actions and nouns for paths. Branch labels should generally stay within three words; explanations belong in the caption region.",
    "Describe AI work plainly: Listening, Finding the paths, or I could not build this yet. Avoid claims that the system sees the future.",
    "Write errors beside the affected control, explain the recovery action, and preserve the person’s input.",
    "Use may.be in lowercase. In the cinematic wordmark, the dot is a functional object and must be mathematically centred before may and be are positioned around it.",
]:
    add_bullet(doc, item)

doc.add_heading("2.11 Figma and code token handoff", level=2)
add_table(doc, ["Figma collection", "Example names", "Code mapping"], [
    ("Color", "background/canvas, text/primary, action/primary", "Semantic CSS custom properties or SwiftUI Color assets"),
    ("Typography", "title/large, body/default, caption/default", "Text styles mapped to Dynamic Type or CSS clamp rules"),
    ("Spacing", "space/1 through space/12", "4 point base scale constants"),
    ("Radius", "radius/control, radius/field, radius/sheet, radius/full", "12, 18, 24, and 999 values"),
    ("Motion", "duration/spatial, easing/cinematic, spring/direct", "Shared motion constants; reduced-motion alternatives beside each token"),
    ("Components", "Mic Orb, Composer, Decision Seed, Branch Node, Caption, Undo", "Variants mirror code states and analytics events"),
], [1.55, 2.55, 2.55])
doc.add_paragraph(
    "Component variants must use identical state names in Figma and code: idle, focused, pressed, listening, processing, current, selected, rejected, loading, error, and disabled where applicable. Prototype flows should demonstrate the complete journey at the 390 by 844 reference size plus one compact and one large Dynamic Type case."
)

# Journey map
doc.add_heading("3 End to end journey", level=1)
doc.add_paragraph("The journey below is the canonical product flow. Each stage owns the screen and advances only when its completion condition is satisfied.")

journey_rows = [
    ("1 Invitation", "See a calm entry point and understand that speech or text can be used.", "Large microphone, short greeting, compact text field."),
    ("2 Capture", "Speak or type the situation in natural language.", "Interim transcript or typed text remains visible and editable."),
    ("3 Listening", "Know the system received the story and is working.", "Breathing sound orb with minimal listening copy."),
    ("4 Meaning", "Recognise the important ideas extracted from the story.", "Grounded keyword bubbles rise from the bowl and dissolve rhythmically."),
    ("5 Passage", "Move emotionally from telling the story to seeing paths.", "Short transition sentence fades, followed by a deliberate empty pause."),
    ("6 Transformation", "See may.be become the decision itself.", "Wordmark appears; may and be fade outward; the centred dot remains."),
    ("7 Decision seed", "Confirm the distilled question before branching.", "The dot expands in place and then reveals the decision label."),
    ("8 Branching", "See all distinct credible first level choices.", "Organic paths grow from the seed; branch nodes arrive as paths reach them."),
    ("9 Exploration", "Understand and follow one path.", "Tap or focus reveals concise detail; selection grows the same continuous tree."),
    ("10 Recovery", "Return to an earlier fork or start again.", "Undo restores the previous branch state; brand action resets the journey."),
]
add_table(doc, ["Stage", "User goal", "Experience response"], journey_rows, [1.15, 2.25, 2.95])

doc.add_heading("3.1 Journey completion criteria", level=2)
for item in [
    "The first decision graph appears only after valid structured analysis has been received.",
    "The user can inspect every visible branch without selecting it.",
    "Selecting a branch preserves spatial continuity and loads the next level without replacing the scene.",
    "Undo restores the exact previous fork, alternatives, and camera position.",
    "The user can always restart from the brand control without reloading the browser page.",
]:
    add_bullet(doc, item)

# Stories
doc.add_heading("4 Epic A Capture the situation", level=1)
add_story(doc, "US 01", "Start with voice", "As a person who is unsure how to frame my problem, I want to speak naturally so that I can begin without composing a formal prompt.", [
    "Given the invitation screen, when I tap the microphone, then recording starts and the microphone visibly enters a listening state.",
    "Given speech recognition is available, when I speak, then interim and final transcript text appears in the text field.",
    "When recording ends with a non-empty final transcript, then the journey starts automatically using that transcript.",
    "When I tap the microphone while recording, then recording stops cleanly.",
], "The initial speech locale is en-IN. A native app may replace browser speech recognition with platform ASR without changing the story contract.")

add_story(doc, "US 02", "Use text instead", "As a person who cannot or does not want to speak, I want a compact text input so that I can use the complete experience privately.", [
    "Given the invitation screen, when I type, then the field preserves the full text while visually truncating only when space requires it.",
    "When I press Enter or tap Send with valid input, then analysis starts once.",
    "The text path provides the same interpretation and graph journey as voice input.",
    "The input remains editable until analysis begins.",
])

add_story(doc, "US 03", "Recover from capture problems", "As a person whose browser or permissions prevent voice capture, I want a clear fallback so that I am never trapped at the first screen.", [
    "Given speech recognition is unsupported, when I tap the microphone, then the app explains that I can use Chrome or Edge or type below.",
    "Given microphone access is denied, when the browser reports not allowed, then the app explains how to enable access and keeps text input usable.",
    "Given speech cannot be understood, then the app offers a calm retry message without discarding typed content.",
    "Error messages are announced to assistive technology and do not cover the primary controls.",
])

add_story(doc, "US 04", "Restart the journey", "As a returning explorer, I want to reset from the may.be brand so that I can describe a different decision without refreshing the page.", [
    "When I activate the brand control, then active timers, animations, analysis state, branch history, captions, and errors are cleared.",
    "The invitation scene returns with the input available for a new situation.",
    "No stale branch nodes or keyword bubbles remain after reset.",
])

doc.add_heading("5 Epic B Turn the story into meaning", level=1)
add_story(doc, "US 05", "See that analysis is active", "As a person waiting for interpretation, I want calm feedback so that processing time feels intentional rather than broken.", [
    "After valid submission, the invitation screen transitions to one breathing orb and minimal listening copy.",
    "The listening state remains visible for at least the minimum cinematic hold while analysis runs.",
    "If analysis completes early, the experience does not flash abruptly to the next scene.",
    "If analysis fails, the app returns to the invitation and shows a retryable error.",
])

add_story(doc, "US 06", "Recognise my own words", "As a person who shared a nuanced situation, I want the extracted words to reflect what I said so that I can trust the interpretation.", [
    "The analysis returns between one and eight concise keywords.",
    "Keywords prioritise explicit options, people, places, constraints, values, and stakes present in the input.",
    "Generic filler such as values, change, risk, or timing is not used when more specific user language is available.",
    "For a direct alternative such as Android or iOS, the displayed keywords are Android and iOS.",
])

add_story(doc, "US 07", "Experience the keyword rhythm", "As a person watching the app make sense of my story, I want keywords to emerge and disappear naturally so that the transition feels alive rather than mechanical.", [
    "Each keyword rises individually from the bowl with a slight stagger.",
    "Each bubble enlarges, overshoots, settles, pauses, and then dissolves with lift and blur.",
    "The sequence is rhythmic and readable; bubbles do not overlap important content or leave the viewport.",
    "The scene advances only after all keyword dissolving is complete.",
], "Target cadence is approximately 430 ms between arrivals with a breathing ease curve of cubic-bezier(.22, 1, .36, 1).")

add_story(doc, "US 08", "Receive a structured decision", "As the experience renderer, I need validated structured analysis so that the interface can create meaning without parsing prose.", [
    "The response contains keywords, framing, factors, a root decision node, decision nodes, and edges.",
    "Every node has a short label, concise summary, grounding kind, factor references, and a human-input flag.",
    "Labels remain compact enough for circular nodes and summaries remain concise enough for the fixed caption area.",
    "Malformed provider output is rejected before it reaches the frontend.",
])

add_story(doc, "US 09", "Respect direct either or questions", "As a person comparing two named options, I want those exact options to become the first branches so that the AI does not replace my decision with generic advice.", [
    "Given input matching A or B, A vs B, or A versus B, the root label is A or B.",
    "The first branches use the two named alternatives in the same order.",
    "Known product names preserve expected casing, including iOS and iPhone.",
    "The AI may add detail beneath each branch but does not introduce unrelated root alternatives.",
])

doc.add_heading("6 Epic C Transform meaning into a decision", level=1)
add_story(doc, "US 10", "Cross the passage", "As a person who has finished telling the story, I want a short acknowledgement before paths appear so that the experience has emotional pacing.", [
    "After keywords dissolve, the app shows a short acknowledgement and path-finding sentence.",
    "The copy fades with blur and slight upward movement rather than disappearing instantly.",
    "A brief blank pause separates the passage from the logo appearance.",
    "No additional instructions or analysis text crowd the scene.",
])

add_story(doc, "US 11", "See the centred logo transformation", "As a person entering the decision space, I want the may.be mark to transform into the decision seed so that the brand and interaction feel like one continuous idea.", [
    "The cinematic may.be wordmark appears with the dot fixed at the mathematical centre of the screen.",
    "The may and be text align around the centred dot and do not move the dot because of unequal word widths.",
    "The may and be text fade outward while the same dot stays in place.",
    "The small header logo remains conventional and is not changed by the cinematic wordmark treatment.",
])

add_story(doc, "US 12", "Understand the decision seed", "As a person, I want the final decision question to appear inside the enlarged dot so that I know what the tree is branching from.", [
    "The dot visibly grows from its original position into the current decision node.",
    "The dot does not juggle or translate while growing.",
    "Decision text appears only after the node reaches a usable size.",
    "The decision label is concise, readable, and derived from the validated root node.",
])

doc.add_heading("7 Epic D Explore the living decision tree", level=1)
add_story(doc, "US 13", "See every credible first path", "As a person considering a decision, I want all distinct credible alternatives to grow from the decision seed so that I can see more than a forced binary choice.", [
    "The renderer supports two to six branches returned by the analysis contract.",
    "The trunk and curved branch paths grow before their destination nodes appear.",
    "Choice nodes arrive with a slight stagger as their paths reach them.",
    "No extra branches are invented merely to fill visual space.",
    "The root, paths, and choices share one world coordinate system.",
])

add_story(doc, "US 14", "Inspect a path before following it", "As a cautious explorer, I want a concise explanation for each branch so that I can understand what it represents before selecting it.", [
    "Hovering, focusing, or touching a choice reveals its summary in a fixed centred caption region.",
    "The explanation never extends outside the mobile viewport.",
    "The node label remains short while the caption carries the detail.",
    "Moving away or losing focus restores the neutral interaction hint.",
])

add_story(doc, "US 15", "Follow one branch", "As a person exploring an alternative, I want the selected branch to become the next decision point so that I can follow its consequences.", [
    "On selection, unselected siblings shrink, blur, move outward, and fade instead of vanishing abruptly.",
    "The chosen node changes into the current node without being replaced by a new visual object.",
    "The camera pans and scales to centre the selected world position.",
    "Existing connective geometry remains coherent throughout the shift.",
    "Input is temporarily locked during the transition to prevent duplicate selection.",
])

add_story(doc, "US 16", "Continue into nearby consequences", "As a person following a path, I want its next credible options ready together so that local AI latency does not break the flow.", [
    "The system prefetches expansions for up to three nearby branches after the current level is received.",
    "Selecting a prefetched branch uses the cached expansion rather than launching a duplicate request.",
    "The next level presents up to three compact child labels for smooth mobile exploration.",
    "If expansion is still running, the selected-node transition provides enough cinematic time to mask ordinary latency.",
    "If expansion fails, the current tree remains intact and the person receives a retryable message.",
])

add_story(doc, "US 17", "Fit larger forks", "As a person whose decision has many credible alternatives, I want the graph to adapt so that choices remain readable and tappable.", [
    "For four branches, world-space spacing increases and the camera zooms out modestly.",
    "For five or more branches, spacing and camera scale increase again.",
    "Labels are not compressed into unreadable text to preserve the default camera framing.",
    "Node hit areas remain at least 44 by 44 CSS pixels at the active zoom.",
])

doc.add_heading("8 Epic E Reverse and trust the exploration", level=1)
add_story(doc, "US 18", "Undo one decision", "As a person who wants to compare another route, I want to undo one branch selection so that exploration never feels like commitment.", [
    "Undo becomes visible only after at least one branch has been followed.",
    "Activating Undo removes only the current child level and restores the previous alternatives.",
    "The previous current node, branch paths, caption state, and camera coordinates are restored seamlessly.",
    "Undo does not refresh the page, restart analysis, or reconstruct the entire graph.",
    "Undo hides again when history becomes empty.",
])

add_story(doc, "US 19", "Know what the AI knows", "As a person making a meaningful choice, I want facts, evidence, assumptions, and human unknowns kept distinct so that I do not mistake generated possibilities for truth.", [
    "Every factor and node is classified as user fact, evidence, assumption, or human unknown.",
    "The product does not display invented probabilities or imply certainty.",
    "Assumptions and human unknowns are editable in a future slice without changing the graph contract.",
    "The product language describes paths as plausible and exploratory, not predicted outcomes.",
])

add_story(doc, "US 20", "Use the journey accessibly", "As a person using touch, keyboard, screen reader, or reduced motion, I want equivalent access so that the experience does not depend on animation alone.", [
    "Every branch is exposed as an operable button with an accessible label containing its short label and summary.",
    "Enter and Space select a focused branch.",
    "Scene changes and errors are announced politely without repeatedly interrupting the user.",
    "Reduced motion shortens staged animations while preserving the order and meaning of each state change.",
    "Text, paths, controls, and focus indicators meet mobile contrast and visibility requirements.",
])

doc.add_heading("9 AI and data contract", level=1)
doc.add_paragraph("The backend returns meaning; the frontend owns choreography. Provider output is normalised and validated before rendering.")
add_table(doc, ["Object", "Required behavior", "Constraint"], [
    ("Input", "Natural-language situation and locale", "12 to 6000 trimmed characters"),
    ("Keywords", "Specific concepts grounded in the input", "1 to 8 items; 48 characters each"),
    ("Factors", "Facts, evidence, assumptions, and unknowns", "At least one factor; editable by default"),
    ("Root node", "Distilled decision question and user context", "One root referenced by rootNodeId"),
    ("Choice node", "Compact path label with concise summary", "Label up to 48; summary up to 320"),
    ("Expansion", "Children of a selected path", "1 to 6 contractually; UI normalises first nearby set"),
], [1.15, 3.05, 2.15])

doc.add_heading("9.1 Grounding model", level=2)
add_table(doc, ["Grounding kind", "Meaning", "Product treatment"], [
    ("user_fact", "Directly supplied by the person", "Treat as current context; allow correction"),
    ("evidence", "Obtained from a named source", "Show provenance when evidence retrieval is added"),
    ("assumption", "Modelled possibility", "Keep editable and never present as fact"),
    ("human_unknown", "A value or missing answer only the person can supply", "Return the question to the person"),
], [1.35, 2.45, 2.55])

doc.add_heading("9.2 Provider behavior", level=2)
for item in [
    "Ollama and LM Studio adapters must return the shared structured schema.",
    "Direct either-or input is normalised into the exact named alternatives.",
    "General decisions are capped at a small set of distinct root branches.",
    "Expanded labels are compacted to approximately three words and 22 characters for the mobile graph.",
    "Analysis and expansion requests are cached per decision and node during the active journey.",
]:
    add_bullet(doc, item)

doc.add_heading("10 Motion and spatial acceptance", level=1)
add_table(doc, ["Moment", "Required motion", "Failure to avoid"], [
    ("Keyword arrival", "Rise, overshoot, settle, pause", "Random floating or simultaneous popping"),
    ("Keyword exit", "Sequential lift, growth, blur, fade", "Abrupt disappearance"),
    ("Passage exit", "Slow fade and slight rise", "Immediate logo collision"),
    ("Logo to seed", "Pinned dot expands in place", "Dot moving between coordinate systems"),
    ("Branch reveal", "Path draws before node arrives", "Nodes popping without connection"),
    ("Selection", "Rejected siblings dissolve; chosen node persists", "Whole scene replacement"),
    ("Camera", "Pan and zoom the shared SVG viewBox", "Moving nodes separately from paths"),
    ("Undo", "Reverse camera and restore prior fork", "Refresh or graph reconstruction"),
], [1.25, 3.0, 2.1])

doc.add_heading("10.1 Timing guidance", level=2)
doc.add_paragraph("Spatial movement uses cubic-bezier(.22, 1, .36, 1). Timings may be tuned through usability testing, but the relative order is mandatory: explain, transform, reveal, then enable interaction. The interface must never enable a branch before its node is visible and stable.")

doc.add_heading("11 Failure and edge cases", level=1)
add_table(doc, ["Condition", "Expected response"], [
    ("Empty or very short input", "Keep the invitation active and ask for enough context without starting analysis."),
    ("Unsupported speech recognition", "Explain the limitation and preserve the fully functional text path."),
    ("Microphone permission denied", "Show a specific permission message and leave text input available."),
    ("Analysis service unavailable", "Return to invitation with a calm retry message; do not fake a graph."),
    ("Provider returns malformed JSON", "Reject it during schema validation and handle as analysis failure."),
    ("No meaningful branches", "Use a deterministic local fallback or ask for clarification; never render an empty graph."),
    ("Slow expansion", "Use prefetch and transition time; keep the selected node and existing graph stable."),
    ("Expansion failure", "Unlock input, retain the current tree, and show a retryable path error."),
    ("Rapid repeated taps", "Ignore additional choices while the transition lock is active."),
    ("Viewport is narrow", "Keep the phone canvas clipped and move detail into the fixed caption area."),
], [2.15, 4.2])

doc.add_heading("12 Definition of done for MVP", level=1)
for item in [
    "A person can complete the journey using either voice or text.",
    "The rendered keywords are specific to the submitted story.",
    "Direct comparisons produce the exact named first branches.",
    "The cinematic passage, logo, centred-dot growth, and decision seed occur in the intended order.",
    "The graph supports a variable number of branches without overlap or viewport escape.",
    "Branch details remain concise and contained in the fixed caption region.",
    "Selecting a path preserves one continuous spatial tree and expands the next level.",
    "Undo restores the previous fork without a page refresh or re-analysis.",
    "AI output is validated and grounded using the shared contract.",
    "Keyboard, screen-reader, touch, and reduced-motion paths are testable.",
    "The light, dark, Increased Contrast, Reduce Transparency, and Dynamic Type variants are implemented from semantic design tokens.",
    "All interactive targets are at least 44 by 44 points and remain reachable inside device safe areas.",
    "Figma component variants and frontend state names match the shared component contract.",
    "All primary error states preserve a route to retry or use text input.",
]:
    add_bullet(doc, item)

doc.add_heading("12.1 Product acceptance scenarios", level=2)
scenarios = [
    ("Career relocation", "A better job in Bangalore means leaving my family. Should I go?", "Specific keywords; root Move to Bangalore; credible Go, Stay, Negotiate paths."),
    ("Direct platform choice", "Android or iOS?", "Keywords Android and iOS; root Android or iOS; branches Android and iOS."),
    ("Uncertain action", "Should I launch now or learn more first?", "Named alternatives preserved; expanded path offers small reversible next steps."),
    ("Voice unavailable", "Browser without speech recognition", "Actionable fallback message; full typed journey remains available."),
    ("Large fork", "AI returns five distinct options", "Wider world positions and zoomed-out camera; labels and hit areas remain usable."),
]
add_table(doc, ["Scenario", "Input or condition", "Expected outcome"], scenarios, [1.4, 2.5, 2.45])

doc.add_heading("12.2 Release boundaries", level=2)
doc.add_paragraph("The MVP is complete when the journey is coherent and reversible. It is not required to recommend one optimal choice, assign probabilities, store decisions across devices, retrieve external evidence, or generate a dashboard. Those capabilities must not be simulated with unsupported claims.")

doc.add_heading("13 Story traceability", level=1)
trace = [
    ("Capture", "US 01 to US 04", "Invitation, voice, text, errors, reset"),
    ("Interpretation", "US 05 to US 09", "Listening, keywords, structured response, direct alternatives"),
    ("Transformation", "US 10 to US 12", "Passage, logo, centred dot, decision seed"),
    ("Exploration", "US 13 to US 17", "Variable branches, detail, selection, expansion, camera"),
    ("Trust and recovery", "US 18 to US 20", "Undo, grounding, accessibility"),
]
add_table(doc, ["Epic", "Stories", "Primary verification area"], trace, [1.55, 1.55, 3.25])

doc.add_heading("14 Design references", level=1)
doc.add_paragraph(
    "The design system uses Apple Human Interface Guidelines as the platform reference current at the time of this specification. Product-specific tokens remain owned by may.be and should be revalidated when platform guidance changes."
)
for label, url in [
    ("Apple Human Interface Guidelines Accessibility", "https://developer.apple.com/design/human-interface-guidelines/accessibility"),
    ("Apple Human Interface Guidelines Color", "https://developer.apple.com/design/human-interface-guidelines/color"),
    ("Apple Human Interface Guidelines Layout", "https://developer.apple.com/design/human-interface-guidelines/layout"),
    ("Apple Human Interface Guidelines Materials", "https://developer.apple.com/design/human-interface-guidelines/materials"),
    ("Apple Human Interface Guidelines Motion", "https://developer.apple.com/design/human-interface-guidelines/motion"),
    ("Apple Human Interface Guidelines Typography", "https://developer.apple.com/design/human-interface-guidelines/typography"),
    ("Apple Human Interface Guidelines Writing", "https://developer.apple.com/design/human-interface-guidelines/writing"),
]:
    add_label_paragraph(doc, label, url)

# Footer
for sec in doc.sections:
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("may.be   Product and Design Specification")
    set_font(r, size=8, color=MUTED)

# Document metadata
doc.core_properties.title = "may.be Product and Design Specification"
doc.core_properties.subject = "User journey, Apple aligned design system, interaction model, and engineering contract"
doc.core_properties.author = "may.be product team"
doc.core_properties.keywords = "may.be, product specification, design system, Apple HIG, user stories, decision exploration, mobile, AI, accessibility"

doc.save(OUT)
print(OUT)
